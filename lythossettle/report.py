"""
Calculation report for Lythos Settle.

The report is assembled as HTML (inputs, stresses, settlement at each point
and in each layer, consolidation time, checks, figures, warnings, method
notes and, if one was run, the study) and exported as
  * PDF   — through reportlab (see `lythossettle.pdf`)
  * HTML  — a single self-contained file (figures embedded as base64)
  * DOCX  — optional, needs `python-docx`

There is one assembly, `build_html()`, so all three formats say the same thing.
"""

from __future__ import annotations

import base64
import datetime
import html
import io
import math
from typing import Any, Dict, List, Optional

from matplotlib.figure import Figure

from . import study_plots
from .config import APP_NAME, APP_VERSION
from .i18n import TRANSLATIONS, warning_text
from .plotting import PLOT_KEYS, Plotter
from .study import OUTPUTS
from .summary import distortion_text, duration

TEXTS = {
    "en": {
        "title": f"{APP_NAME} — Calculation Report",
        "date": "Date", "analyst": "Analyst", "software": "Software",
        "sec_inputs": "1. Input data", "sec_found": "1.1 Foundation and loading",
        "sec_soil": "1.2 Soil profile", "sec_opts": "1.3 Analysis options and criteria",
        "sec_stress": "2. Stresses", "sec_stress_layers": "2.1 Stresses at the middle of each layer",
        "sec_settle": "3. Settlement", "sec_points": "3.1 Settlement at the evaluation points",
        "sec_layers": "3.2 Settlement by layer ({point})", "sec_schm": "3.3 Schmertmann factors",
        "sec_time": "3.4 Consolidation time", "sec_checks": "3.5 Checks",
        "sec_figs": "4. Figures", "sec_warn": "5. Warnings", "sec_notes": "6. Method notes",
        "parameter": "Parameter", "value": "Value", "unit": "Unit",
        "shape": "Shape", "B": "Width B", "L": "Length L", "D": "Diameter D",
        "Df": "Foundation depth Df", "q": "Bearing pressure q (gross)",
        "net": "Net pressure (overburden deducted)", "zw": "Water table depth",
        "gw": "Unit weight of water γw", "yes": "yes", "no": "no",
        "layer": "Layer", "type": "Type", "top": "Top", "bottom": "Bottom",
        "clay_params": "Consolidation parameters of the cohesive layers",
        "stress_method": "Stress distribution", "imm_method": "Immediate settlement",
        "rigidity": "Foundation rigidity", "sublayer": "Sublayer thickness",
        "emb_crest": "Crest width", "emb_height": "Height H", "emb_left": "Left slope angle",
        "emb_right": "Right slope angle", "emb_base": "Width at the base",
        "emb_gamma": "Unit weight of the fill γ", "emb_q": "Fill load q = γ·H (under the crest)",
        "note_emb": "Embankment: a long trapezoidal fill on the ground surface (plane strain). Δσ "
                    "is exact for the piecewise-linear strip load (Flamant's line load integrated "
                    "over each segment); the elastic settlement superposes uniform strips, the "
                    "crest as one and each slope as 16 slices. The fill's own compression, "
                    "undrained lateral spreading and stability are not included.",
        "depth_ratio": "Influence depth criterion Δσ / σ'v0", "life": "Design life",
        "creep": "Schmertmann creep factor C2", "s_allow": "Allowable total settlement",
        "d_allow": "Allowable angular distortion",
        "sigma_base": "Total overburden at the base σv0", "q_net": "Net pressure q_net",
        "z_limit": "Influence depth (below ground)",
        "z_mid": "Mid-depth", "sigma_eff": "σ'v0", "dsigma": "Δσ", "sigma_p": "σ'p",
        "point": "Point", "coords": "Position", "immediate": "Immediate",
        "consolidation": "Consolidation", "secondary": "Secondary", "total": "Total",
        "method": "Method", "h_dr": "Drainage path H_dr", "cv": "cv", "t50": "t50", "t90": "t90",
        "at_life": "Settlement after {life:g} years (governing point)",
        "check": "Check", "actual": "Actual", "allowable": "Allowable", "status": "Status",
        "c_total": "Total settlement", "c_dist": "Angular distortion (centre – edge)",
        "ok": "OK", "notok": "NOT OK", "na": "n/a",
        "sec_study": "7. Parametric / reliability study", "sec_study_vars": "7.1 Variables",
        "sec_study_stats": "7.2 Statistics of the outputs",
        "sec_study_rel": "7.3 Probability of exceeding the criteria",
        "sec_study_sens": "7.4 Sensitivities (Spearman ρ)", "sec_study_figs": "7.5 Figures",
        "study_method": "Sampling method", "study_n": "Samples", "study_ok": "Successful analyses",
        "variable": "Variable", "mode": "Mode", "range": "Range", "dist": "Distribution",
        "mean": "Mean", "cov": "CoV", "min": "Min", "max": "Max", "output": "Output",
        "std": "Std", "criterion": "Criterion", "n_fail": "Exceeded", "pf": "P",
        "pf_ci": "95 % CI", "beta_idx": "β",
        "notes": [
            "In-situ stresses from the unit weights, γ above and γsat below the water table; "
            "σ'p = OCR·σ'v0 in the cohesive layers.",
            "Net pressure q_net = q − σv0(Df) when the excavated overburden is deducted. The "
            "stress increase Δσ = q_net·I is from Boussinesq (Newmark's rectangle, the strip "
            "and the circle solutions) or from the 2:1 spread, which is an average over the "
            "loaded width.",
            "Settlement is summed over sublayers down to the depth where Δσ at the centre falls "
            "to the chosen fraction of σ'v0; the base of the profile is incompressible.",
            "Immediate settlement, elastic: Steinbrenner's F1, F2 for a flexible rectangle on a "
            "layer of finite thickness, s = q_net/E·Σ B·[(1 − ν²)F1 + (1 − ν − 2ν²)F2] between "
            "the top and bottom of each layer, superposed for any point. A strip is a long "
            "rectangle, a circle the square of the same area. In a cohesive layer E and ν are "
            "the undrained values.",
            "Immediate settlement, Schmertmann et al. (1978) in the granular layers: "
            "s = C1·C2·q_net·Σ Iz/E·Δz, with C1 = 1 − 0.5·σ'v0/q_net ≥ 0.5, "
            "C2 = 1 + 0.2·log10(t/0.1), Izp = 0.5 + 0.1·√(q_net/σ'vp) and the influence diagram "
            "interpolated between L/B = 1 and L/B ≥ 10. The value is the footing's settlement; "
            "away from the centre it is scaled by the elastic distribution.",
            "Primary consolidation: Δe from Cr up to σ'p and Cc beyond it, "
            "s = Σ Δe/(1 + e0)·Δz, sublayer by sublayer at each point.",
            "Time: Terzaghi's one-dimensional theory for a uniform initial excess pore pressure, "
            "each clay layer draining on its own, H_dr = H/2 (double) or H (single drainage).",
            "Secondary compression from the end of primary consolidation (U = 95 %) to the "
            "design life: s = Cα/(1 + e0)·H·log10(t/t_p); where the final stress stays below "
            "σ'p the rate is reduced to Cα·Cr/Cc (Mesri's Cα/Cc concept).",
            "A rigid foundation settles uniformly by the settlement of its characteristic point "
            "(0.74·B/2 and 0.74·L/2 from the centre; 0.845·R in a circle). The angular "
            "distortion of a flexible foundation is the difference between the centre and the "
            "middle of the long edge over half the width.",
        ],
    },
    "tr": {
        "title": f"{APP_NAME} — Hesap Raporu",
        "date": "Tarih", "analyst": "Hazırlayan", "software": "Yazılım",
        "sec_inputs": "1. Girdi verileri", "sec_found": "1.1 Temel ve yükleme",
        "sec_soil": "1.2 Zemin profili", "sec_opts": "1.3 Analiz seçenekleri ve ölçütler",
        "sec_stress": "2. Gerilmeler", "sec_stress_layers": "2.1 Tabaka ortalarında gerilmeler",
        "sec_settle": "3. Oturma", "sec_points": "3.1 Hesap noktalarında oturma",
        "sec_layers": "3.2 Tabakalara göre oturma ({point})",
        "sec_schm": "3.3 Schmertmann katsayıları",
        "sec_time": "3.4 Konsolidasyon süresi", "sec_checks": "3.5 Kontroller",
        "sec_figs": "4. Şekiller", "sec_warn": "5. Uyarılar", "sec_notes": "6. Yöntem notları",
        "parameter": "Parametre", "value": "Değer", "unit": "Birim",
        "shape": "Şekil", "B": "Genişlik B", "L": "Boy L", "D": "Çap D",
        "Df": "Temel derinliği Df", "q": "Taban basıncı q (brüt)",
        "net": "Net basınç (örtü yükü düşülmüş)", "zw": "Su tablası derinliği",
        "gw": "Suyun birim hacim ağırlığı γw", "yes": "evet", "no": "hayır",
        "layer": "Tabaka", "type": "Tür", "top": "Üst", "bottom": "Alt",
        "clay_params": "Kohezyonlu tabakaların konsolidasyon parametreleri",
        "stress_method": "Gerilme dağılımı", "imm_method": "Ani oturma",
        "rigidity": "Temel rijitliği", "sublayer": "Alt tabaka kalınlığı",
        "emb_crest": "Tepe genişliği", "emb_height": "Yükseklik H", "emb_left": "Sol şev açısı",
        "emb_right": "Sağ şev açısı", "emb_base": "Taban genişliği",
        "emb_gamma": "Dolgunun birim hacim ağırlığı γ", "emb_q": "Dolgu yükü q = γ·H (tepe altında)",
        "note_emb": "Dolgu: zemin yüzeyinde uzun, yamuk kesitli dolgu (düzlem şekil değiştirme). "
                    "Δσ parçalı doğrusal şerit yük için kesindir (Flamant çizgisel yükü her parça "
                    "üzerinde integre edilir); elastik oturma düzgün şeritlerin süperpozisyonudur, "
                    "tepe tek şerit, her şev 16 dilim. Dolgunun kendi sıkışması, drenajsız yanal "
                    "yayılma ve stabilite hesaba katılmaz.",
        "depth_ratio": "Etki derinliği ölçütü Δσ / σ'v0", "life": "Tasarım ömrü",
        "creep": "Schmertmann sünme katsayısı C2", "s_allow": "İzin verilen toplam oturma",
        "d_allow": "İzin verilen açısal distorsiyon",
        "sigma_base": "Tabanda toplam örtü yükü σv0", "q_net": "Net basınç q_net",
        "z_limit": "Etki derinliği (zeminden)",
        "z_mid": "Orta derinlik", "sigma_eff": "σ'v0", "dsigma": "Δσ", "sigma_p": "σ'p",
        "point": "Nokta", "coords": "Konum", "immediate": "Ani",
        "consolidation": "Konsolidasyon", "secondary": "İkincil", "total": "Toplam",
        "method": "Yöntem", "h_dr": "Drenaj yolu H_dr", "cv": "cv", "t50": "t50", "t90": "t90",
        "at_life": "{life:g} yıl sonundaki oturma (esas nokta)",
        "check": "Kontrol", "actual": "Oluşan", "allowable": "İzin verilen", "status": "Durum",
        "c_total": "Toplam oturma", "c_dist": "Açısal distorsiyon (merkez – kenar)",
        "ok": "UYGUN", "notok": "UYGUN DEĞİL", "na": "—",
        "sec_study": "7. Parametrik / güvenilirlik çalışması", "sec_study_vars": "7.1 Değişkenler",
        "sec_study_stats": "7.2 Çıktıların istatistikleri",
        "sec_study_rel": "7.3 Ölçütlerin aşılma olasılığı",
        "sec_study_sens": "7.4 Duyarlılıklar (Spearman ρ)", "sec_study_figs": "7.5 Şekiller",
        "study_method": "Örnekleme yöntemi", "study_n": "Örnek sayısı",
        "study_ok": "Başarılı analiz", "variable": "Değişken", "mode": "Mod", "range": "Aralık",
        "dist": "Dağılım", "mean": "Ortalama", "cov": "CoV", "min": "Min", "max": "Maks",
        "output": "Çıktı", "std": "Std", "criterion": "Ölçüt", "n_fail": "Aşan", "pf": "P",
        "pf_ci": "%95 GA", "beta_idx": "β",
        "notes": [
            "Yerinde gerilmeler birim hacim ağırlıklarından; su tablası üstünde γ, altında γdoy. "
            "Kohezyonlu tabakalarda σ'p = OCR·σ'v0.",
            "Kazılan örtü yükü düşüldüğünde net basınç q_net = q − σv0(Df). Gerilme artışı "
            "Δσ = q_net·I; Boussinesq (Newmark dikdörtgen, şerit ve daire çözümleri) ya da "
            "yüklenen genişlik üzerinde ortalama veren 2:1 yayılma ile.",
            "Oturma, merkezdeki Δσ'nın σ'v0'ın seçilen oranına düştüğü derinliğe kadar alt "
            "tabakalar boyunca toplanır; profil tabanı sıkışmaz kabul edilir.",
            "Elastik ani oturma: sonlu kalınlıklı tabaka üzerindeki esnek dikdörtgen için "
            "Steinbrenner F1, F2; her tabakanın üstü ile altı arasında "
            "s = q_net/E·Σ B·[(1 − ν²)F1 + (1 − ν − 2ν²)F2], herhangi bir nokta için "
            "süperpozisyon. Şerit uzun bir dikdörtgen, daire eşit alanlı kare olarak alınır. "
            "Kohezyonlu tabakada E ve ν drenajsız değerlerdir.",
            "Granüler tabakalarda Schmertmann vd. (1978) ani oturma: s = C1·C2·q_net·Σ Iz/E·Δz; "
            "C1 = 1 − 0.5·σ'v0/q_net ≥ 0.5, C2 = 1 + 0.2·log10(t/0.1), "
            "Izp = 0.5 + 0.1·√(q_net/σ'vp); etki diyagramı L/B = 1 ile L/B ≥ 10 arasında "
            "enterpole edilir. Değer temelin oturmasıdır; merkez dışında elastik dağılımla "
            "ölçeklenir.",
            "Birincil konsolidasyon: σ'p'ye kadar Cr, ötesinde Cc ile Δe; "
            "s = Σ Δe/(1 + e0)·Δz, her noktada alt tabaka alt tabaka.",
            "Zaman: düzgün başlangıç boşluk suyu basıncı için Terzaghi tek boyutlu teorisi; her "
            "kil tabakası kendi başına drene olur, H_dr = H/2 (çift) veya H (tek yönlü drenaj).",
            "İkincil sıkışma birincil konsolidasyonun sonundan (U = %95) tasarım ömrüne kadar: "
            "s = Cα/(1 + e0)·H·log10(t/t_p); son gerilmenin σ'p'nin altında kaldığı yerde hız "
            "Cα·Cr/Cc'ye indirilir (Mesri Cα/Cc yaklaşımı).",
            "Rijit temel, karakteristik noktasının oturması kadar düzgün oturur (merkezden "
            "0.74·B/2 ve 0.74·L/2; dairede 0.845·R). Esnek temelde açısal distorsiyon, merkez "
            "ile uzun kenar ortası arasındaki farkın yarım genişliğe oranıdır.",
        ],
    },
}

#: The interface's palette: warm paper, ink, terracotta; serif headings
_CSS = """
body { font-family: system-ui, -apple-system, 'Segoe UI', 'Helvetica Neue', Arial, sans-serif;
       font-size: 10pt; color: #141413; }
h1, h2, h3 { font-family: 'Tiempos Text', 'Source Serif 4', 'Iowan Old Style', Palatino,
             Georgia, 'DejaVu Serif', serif; font-weight: 500; }
h1 { font-size: 20pt; color: #141413; margin-bottom: 2px; }
h2 { font-size: 14pt; color: #C6613F; border-bottom: 1px solid #E3E0D5; padding-bottom: 2px;
     margin-top: 20px; }
h3 { font-size: 11.5pt; color: #141413; margin-top: 12px; }
table { border-collapse: collapse; margin: 4px 0 8px 0; }
th { background: #F0EEE6; text-align: left; padding: 3px 6px; border: 1px solid #E3E0D5;
     font-size: 9pt; }
td { padding: 3px 6px; border: 1px solid #E3E0D5; font-size: 9pt; }
.ok { color: #3F7F4F; font-weight: bold; } .bad { color: #B0413E; font-weight: bold; }
.meta { color: #73726C; } .note { color: #73726C; font-size: 9pt; }
"""


def _esc(x) -> str:
    return html.escape(str(x))


def _f(x, nd=2) -> str:
    if x is None or (isinstance(x, float) and not math.isfinite(x)):
        return "—"
    return f"{x:,.{nd}f}"


def _status(T, s: str) -> str:
    if s == "N/A":
        return T["na"]
    return (f'<span class="ok">{T["ok"]}</span>' if s == "OK"
            else f'<span class="bad">{T["notok"]}</span>')


def _table(headers: List[str], rows: List[List[Any]], widths: Optional[List[int]] = None) -> str:
    out = ["<table width='100%'>"]
    if widths:
        cells = "".join(f"<th width='{w}%'>{_esc(h)}</th>" for h, w in zip(headers, widths))
    else:
        cells = "".join(f"<th>{_esc(h)}</th>" for h in headers)
    out.append("<tr>" + cells + "</tr>")
    for r in rows:
        out.append("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>")
    out.append("</table>")
    return "\n".join(out)


def _kv_table(T, rows: List[List[Any]]) -> str:
    return _table([T["parameter"], T["value"], T["unit"]], rows, [50, 36, 14])


# ----------------------------------------------------------------------
def _png(fig: Figure, dpi: int) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight")
    return buf.getvalue()


def render_figures(analysis, lang: str, keys=PLOT_KEYS, dpi: int = 130) -> Dict[str, bytes]:
    """The report figures as PNG bytes (off-screen, light theme)."""
    plotter = Plotter(analysis, lang, "light")
    out = {}
    for key in keys:
        fig = Figure(figsize=(10, 7), dpi=dpi)
        plotter.draw(key, fig)
        out[key] = _png(fig, dpi)
    return out


def render_study_figures(study, lang: str, dpi: int = 110) -> Dict[str, bytes]:
    """Study figures keyed 'study_<view>' (only the views that apply)."""
    if study is None or not study.rows:
        return {}
    L = TRANSLATIONS[lang]
    views = ["oat", "hist"] if study.method == "oat" else ["hist", "scatter", "tornado"]
    draw = {"oat": study_plots.plot_oat, "hist": study_plots.plot_hist,
            "scatter": study_plots.plot_scatter, "tornado": study_plots.plot_tornado}
    out = {}
    for view in views:
        fig = Figure(figsize=(10, 7), dpi=dpi)
        draw[view](fig, study, L, "s_total")
        out[f"study_{view}"] = _png(fig, dpi)
    return out


def _study_section(T, L, study, figures, img_src) -> List[str]:
    parts = [f"<h2 style='page-break-before:always'>{T['sec_study']}</h2>"]
    s = study.summary
    parts.append(_kv_table(T, [
        [T["study_method"], L[f"method_{study.method}"], ""],
        [T["study_n"], str(s.get("n_total", 0)), ""],
        [T["study_ok"], str(s.get("n_ok", 0)), ""],
    ]))
    parts.append(f"<h3>{T['sec_study_vars']}</h3>")
    rows = []
    for v in study.variables:
        if v.mode == "range":
            rows.append([_esc(v.label), T["range"], _f(v.vmin, 3), _f(v.vmax, 3), "—", "—", "—",
                         v.n_points])
        else:
            rows.append([_esc(v.label), T["dist"], "—", "—", L[f"dist_{v.dist}"], _f(v.mean, 3),
                         _f(v.cov, 3), "—"])
    parts.append(_table([T["variable"], T["mode"], T["min"], T["max"], T["dist"], T["mean"],
                         T["cov"], "n"], rows, [28, 12, 10, 10, 14, 10, 8, 8]))
    parts.append(f"<h3>{T['sec_study_stats']}</h3>")
    rows = [[_esc(L.get(f"out_{k}", k)), st["n"], _f(st["mean"]), _f(st["std"]), _f(st["p5"]),
             _f(st["p50"]), _f(st["p95"])]
            for k, _ in OUTPUTS if k in s.get("stats", {}) for st in [s["stats"][k]]]
    parts.append(_table([T["output"], "n", T["mean"], T["std"], "P5", "P50", "P95"],
                        rows, [34, 8, 12, 12, 11, 11, 12]))
    if s.get("reliability"):
        parts.append(f"<h3>{T['sec_study_rel']}</h3>")
        rows = []
        for name, r in s["reliability"].items():
            rows.append([_esc(L[name]), r["n"], r["n_fail"], f"{r['pf']:.3g}",
                         f"[{r['pf_lo']:.2g}, {r['pf_hi']:.2g}]",
                         _esc(study_plots._beta_text(r["beta"], r["n"]))])
        parts.append(_table([T["criterion"], "n", T["n_fail"], T["pf"], T["pf_ci"], T["beta_idx"]],
                            rows, [30, 10, 12, 14, 20, 14]))
    if s.get("spearman"):
        parts.append(f"<h3>{T['sec_study_sens']}</h3>")
        outs = [k for k, _ in OUTPUTS if k in s["spearman"]]
        rows = [[_esc(v.label)] + [f"{s['spearman'][k].get(v.path, float('nan')):+.2f}"
                                   for k in outs] for v in study.variables]
        parts.append(_table([T["variable"]] + [L.get(f"out_{k}", k) for k in outs], rows))
    parts.append(f"<h3>{T['sec_study_figs']}</h3>")
    for key in figures:
        if key.startswith("study_"):
            parts.append(f"<p><img src='{img_src(key)}' width='640'></p>")
    return parts


def build_html(analysis, lang: str, figures: Dict[str, bytes], img_src=None,
               study=None) -> str:
    """
    Builds the report HTML. img_src(key) -> value for the <img src> attribute;
    the default embeds base64 data URIs (the self-contained HTML file). The PDF
    and DOCX writers pass a `fig://<key>` mapper and resolve the keys against
    the figure dictionary themselves.
    """
    T = TEXTS.get(lang, TEXTS["en"])
    L = TRANSLATIONS.get(lang, TRANSLATIONS["en"])
    a, res = analysis, analysis.results
    info = a.config.get("project_info", {})
    if img_src is None:
        def img_src(key):
            return "data:image/png;base64," + base64.b64encode(figures[key]).decode()

    yn = lambda b: T["yes"] if b else T["no"]  # noqa: E731
    parts = [f"<html><head><meta charset='utf-8'><style>{_CSS}</style></head><body>"]
    parts.append(f"<h1>{_esc(info.get('title', ''))}</h1>")
    parts.append(f"<p class='meta'>{_esc(T['title'])}<br>"
                 f"{T['date']}: {datetime.date.today().isoformat()} &nbsp;|&nbsp; "
                 f"{T['analyst']}: {_esc(info.get('analyst', ''))} &nbsp;|&nbsp; "
                 f"{T['software']}: {APP_NAME} v{APP_VERSION}</p>")

    # ---------------- 1. inputs
    parts.append(f"<h2>{T['sec_inputs']}</h2>")
    parts.append(f"<h3>{T['sec_found']}</h3>")
    rows = [[T["shape"], L["shape_" + a.shape], ""]]
    if a.embankment:
        e = a.embankment
        rows += [[T["emb_crest"], _f(e["crest"]), "m"], [T["emb_height"], _f(e["height"]), "m"],
                 [T["emb_left"], f"{_f(e['slope_left'])} (1:{_f(e['run_left'] / e['height'])})",
                  "°"],
                 [T["emb_right"],
                  f"{_f(e['slope_right'])} (1:{_f(e['run_right'] / e['height'])})", "°"],
                 [T["emb_base"], _f(a.B), "m"], [T["emb_gamma"], _f(e["gamma"], 1), "kN/m³"],
                 [T["emb_q"], _f(a.q, 1), "kPa"]]
    else:
        if a.shape == "circle":
            rows.append([T["D"], _f(a.B), "m"])
        else:
            rows.append([T["B"], _f(a.B), "m"])
            if a.shape == "rectangle":
                rows.append([T["L"], _f(a.L), "m"])
        rows += [[T["Df"], _f(a.Df), "m"], [T["q"], _f(a.q, 1), "kPa"],
                 [T["net"], yn(a.net_pressure), ""]]
    rows += [[T["zw"], _f(a.profile.zw), "m"], [T["gw"], _f(a.profile.gw), "kN/m³"]]
    parts.append(_kv_table(T, rows))

    parts.append(f"<h3>{T['sec_soil']}</h3>")
    rows = [[i, _esc(lay["name"]), L["behaviour_" + lay["behaviour"]], _f(lay["top"]),
             _f(lay["bottom"]), _f(lay["gamma"], 1), _f(lay["gamma_sat"], 1),
             _f(lay["E_MPa"], 1), _f(lay["nu"])]
            for i, lay in enumerate(a.profile.layers, 1)]
    parts.append(_table(["#", T["layer"], T["type"], f"{T['top']} (m)", f"{T['bottom']} (m)",
                         "γ (kN/m³)", L["col_gamma_sat"], "E (MPa)", "ν"], rows,
                        [4, 24, 12, 8, 8, 11, 12, 11, 10]))
    clays = [(i, lay) for i, lay in enumerate(a.profile.layers, 1)
             if lay["behaviour"] == "cohesive"]
    if clays:
        parts.append(f"<p class='lead'><b>{_esc(T['clay_params'])}</b></p>")
        rows = [[i, _esc(lay["name"]), _f(lay["Cc"], 3), _f(lay["Cr"], 3), _f(lay["e0"]),
                 _f(lay["OCR"]), _f(lay["cv"]), _f(lay["Calpha"], 4),
                 L["drainage_" + lay["drainage"]]] for i, lay in clays]
        parts.append(_table(["#", T["layer"], "Cc", "Cr", "e0", "OCR", L["col_cv"], "Cα",
                             L["col_drainage"]], rows, [4, 24, 8, 8, 8, 8, 14, 10, 16]))

    parts.append(f"<h3>{T['sec_opts']}</h3>")
    parts.append(_kv_table(T, [
        [T["stress_method"], L["stress_" + a.stress_method], ""],
        [T["imm_method"], L["immediate_" + a.immediate_method], ""],
        [T["rigidity"], L["rigidity_" + a.rigidity], ""],
        [T["sublayer"], _f(a.sublayer), "m"],
        [T["depth_ratio"], _f(a.depth_ratio), ""],
        [T["life"], _f(a.design_life, 1), L["unit_years"].format(v=0).split(" ")[-1]],
        [T["creep"], yn(a.creep), ""],
        [T["s_allow"], _f(a.s_allow, 1) if a.s_allow > 0 else T["na"], "mm"],
        [T["d_allow"], f"1/{a.distortion_allow:.0f}" if a.distortion_allow > 0 else T["na"], ""],
    ]))

    # ---------------- 2. stresses
    parts.append(f"<h2 style='page-break-before:always'>{T['sec_stress']}</h2>")
    if a.embankment:
        rows = [[T["emb_q"], _f(a.q, 1), "kPa"]]
    else:
        rows = [[T["q"], _f(a.q, 1), "kPa"], [T["sigma_base"], _f(res["sigma_base"], 1), "kPa"],
                [T["q_net"], _f(res["q_net"], 1), "kPa"]]
    parts.append(_kv_table(T, rows + [[T["z_limit"], _f(res["z_limit"]), "m"]]))
    parts.append(f"<h3>{T['sec_stress_layers']}</h3>")
    gov = res["governing"]
    rows = []
    for info_row in res["layers"]:
        z = 0.5 * (info_row["top"] + info_row["bottom"])
        sig = float(a.profile.effective_stress(z)[0])
        ds = float(res["q_net"] * a.influence(a.points()[gov], z - a.Df)[0])
        lay = a.profile.layers[info_row["index"]]
        sp = sig * lay["OCR"] if lay["behaviour"] == "cohesive" else None
        rows.append([_esc(info_row["name"]), _f(z), _f(sig, 1), _f(ds, 1), _f(sp, 1)])
    parts.append(_table([T["layer"], f"{T['z_mid']} (m)", f"{T['sigma_eff']} (kPa)",
                         f"{T['dsigma']} — {L['point_' + gov].lower()} (kPa)",
                         f"{T['sigma_p']} (kPa)"], rows))

    # ---------------- 3. settlement
    parts.append(f"<h2>{T['sec_settle']}</h2>")
    parts.append(f"<h3>{T['sec_points']}</h3>")
    rows = []
    for key, p in res["points"].items():
        coords = ", ".join(f"{c:.2f}" for c in p["coords"])
        rows.append([L["point_" + key], f"({coords})", _f(p["immediate"], 1),
                     _f(p["consolidation"], 1), _f(p["secondary"], 1),
                     f"<b>{_f(p['total'], 1)}</b>"])
    parts.append(_table([T["point"], f"{T['coords']} (m)", f"{T['immediate']} (mm)",
                         f"{T['consolidation']} (mm)", f"{T['secondary']} (mm)",
                         f"{T['total']} (mm)"], rows, [24, 16, 15, 15, 15, 15]))

    parts.append(f"<h3>{T['sec_layers'].format(point=L['point_' + gov].lower())}</h3>")
    rows = []
    for info_row in res["layers"]:
        total = info_row["immediate"] + info_row["consolidation"] + info_row["secondary"]
        rows.append([_esc(info_row["name"]), L["method_" + info_row["method"]],
                     _f(info_row["immediate"], 1), _f(info_row["consolidation"], 1),
                     _f(info_row["secondary"], 1), _f(total, 1)])
    parts.append(_table([T["layer"], T["method"], f"{T['immediate']} (mm)",
                         f"{T['consolidation']} (mm)", f"{T['secondary']} (mm)",
                         f"{T['total']} (mm)"], rows, [22, 26, 13, 13, 13, 13]))

    schm = res.get("schmertmann")
    if schm:
        parts.append(f"<h3>{T['sec_schm']}</h3>")
        parts.append(_kv_table(T, [
            ["C1", _f(schm["C1"], 3), ""], ["C2", _f(schm["C2"], 3), ""],
            ["Iz0", _f(schm["Iz0"], 3), ""], ["Izp", _f(schm["Izp"], 3), ""],
            ["z_p", _f(schm["z_peak"]), "m"], ["z_end", _f(schm["z_end"]), "m"],
            ["σ'vp", _f(schm["sigma_peak"], 1), "kPa"],
        ]))

    parts.append(f"<h3>{T['sec_time']}</h3>")
    clays = [row for row in res["layers"] if row["behaviour"] == "cohesive"]
    if clays:
        rows = [[_esc(row["name"]), _f(row["h_dr"]), _f(row["cv"]), duration(L, row["t50"]),
                 duration(L, row["t90"])] for row in clays]
        parts.append(_table([T["layer"], f"{T['h_dr']} (m)", f"{T['cv']} ({L['col_cv'][4:-1]})",
                             T["t50"], T["t90"]], rows))
    parts.append(f"<p>{_esc(T['at_life'].format(life=a.design_life))}: "
                 f"<b>{_f(res['at_life'], 1)} mm</b></p>")

    parts.append(f"<h3>{T['sec_checks']}</h3>")
    check_t, check_d = res["checks"]["total"], res["checks"]["distortion"]
    rows = [[T["c_total"], f"{_f(check_t['actual'], 1)} mm",
             f"{_f(check_t['allowable'], 1)} mm" if check_t["status"] != "N/A" else T["na"],
             _status(T, check_t["status"])],
            [T["c_dist"], distortion_text(check_d["actual"]),
             distortion_text(check_d["allowable"]) if check_d["allowable"] else T["na"],
             _status(T, check_d["status"])]]
    parts.append(_table([T["check"], T["actual"], T["allowable"], T["status"]], rows,
                        [40, 20, 20, 20]))
    if a.embankment:
        parts.append(f"<p class='note'>{_esc(L['res_emb_dist'])}</p>")
    elif a.rigidity == "rigid":
        parts.append(f"<p class='note'>{_esc(L['res_rigid'])}</p>")

    # ---------------- 4. figures
    parts.append(f"<h2 style='page-break-before:always'>{T['sec_figs']}</h2>")
    for key in [k for k in figures if not k.startswith("study_")]:
        parts.append(f"<p class='lead'><b>{_esc(L.get(f'fig_{key}', key))}</b></p>")
        parts.append(f"<p><img src='{img_src(key)}' width='640'></p>")

    # ---------------- 5/6. warnings & notes
    if res.get("warnings"):
        parts.append(f"<h2>{T['sec_warn']}</h2><ul>")
        parts.extend(f"<li>{_esc(warning_text(lang, w))}</li>" for w in res["warnings"])
        parts.append("</ul>")
    parts.append(f"<h2>{T['sec_notes']}</h2><ul class='note'>")
    notes = T["notes"] + ([T["note_emb"]] if a.embankment else [])
    parts.extend(f"<li>{_esc(n)}</li>" for n in notes)
    parts.append("</ul>")
    if study is not None and study.rows:
        parts.extend(_study_section(T, L, study, figures, img_src))
    parts.append("</body></html>")
    return "\n".join(parts)


# ----------------------------------------------------------------------
# exporters
# ----------------------------------------------------------------------
def _all_figures(analysis, lang, study):
    figures = render_figures(analysis, lang)
    figures.update(render_study_figures(study, lang))
    return figures


def export_html(path: str, analysis, lang: str, study=None) -> None:
    figures = _all_figures(analysis, lang, study)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(build_html(analysis, lang, figures, study=study))


def export_pdf(path: str, analysis, lang: str, study=None) -> None:
    """PDF through reportlab; the figures travel as `fig://<key>` references."""
    from . import pdf as pdf_writer

    T = TEXTS.get(lang, TEXTS["en"])
    info = analysis.config.get("project_info", {})
    figures = _all_figures(analysis, lang, study)
    html_text = build_html(analysis, lang, figures, img_src=lambda k: f"fig://{k}", study=study)
    pdf_writer.html_to_pdf(path, html_text, figures,
                           title=info.get("title", T["title"]),
                           author=info.get("analyst", ""),
                           footer=f"{info.get('title', '')} — {APP_NAME} v{APP_VERSION}")


def export_docx(path: str, analysis, lang: str, study=None) -> None:
    """Word report; requires python-docx."""
    try:
        import docx
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("python-docx is not installed (pip install python-docx).") from exc
    import re

    figures = _all_figures(analysis, lang, study)
    html_text = build_html(analysis, lang, figures, img_src=lambda k: f"fig://{k}", study=study)
    d = docx.Document()
    from docx.shared import RGBColor
    d.styles["Normal"].font.size = Pt(10)
    d.styles["Normal"].font.color.rgb = RGBColor(0x14, 0x14, 0x13)
    # the interface's look: serif headings, the title in ink, sections in terracotta
    for name, colour in (("Title", (0x14, 0x14, 0x13)), ("Heading 1", (0xC6, 0x61, 0x3F)),
                         ("Heading 2", (0x14, 0x14, 0x13))):
        style = d.styles[name]
        style.font.name = "Georgia"
        style.font.bold = False
        style.font.color.rgb = RGBColor(*colour)

    # a small HTML -> docx walker (headings, paragraphs, tables, lists, images)
    tokens = re.split(r"(<h1>.*?</h1>|<h2[^>]*>.*?</h2>|<h3>.*?</h3>|<table[^>]*>.*?</table>|"
                      r"<p[^>]*>.*?</p>|<li>.*?</li>)", html_text, flags=re.S)

    def strip(s):
        return html.unescape(re.sub(r"<[^>]+>", "", re.sub(r"<br\s*/?>", "\n", s))).strip()

    for tok in tokens:
        if tok.startswith("<h1>"):
            d.add_heading(strip(tok), level=0)
        elif tok.startswith("<h2"):
            d.add_heading(strip(tok), level=1)
        elif tok.startswith("<h3>"):
            d.add_heading(strip(tok), level=2)
        elif tok.startswith("<li>"):
            d.add_paragraph(strip(tok), style="List Bullet")
        elif tok.startswith("<table"):
            rows = re.findall(r"<tr>(.*?)</tr>", tok, flags=re.S)
            cells = [re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", r, flags=re.S) for r in rows]
            if cells:
                table = d.add_table(rows=len(cells), cols=len(cells[0]))
                table.style = "Light Grid Accent 1"
                for i, row in enumerate(cells):
                    for j, c in enumerate(row):
                        if j < len(table.columns):
                            table.cell(i, j).text = strip(c)
        elif tok.startswith("<p"):
            m = re.search(r"src='fig://([a-z_]+)'", tok)
            if m:
                d.add_picture(io.BytesIO(figures[m.group(1)]), width=Inches(6.3))
            else:
                text = strip(tok)
                if text:
                    d.add_paragraph(text)
    d.save(path)


def export_report(path: str, analysis, lang: str, study=None) -> None:
    ext = path.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        export_pdf(path, analysis, lang, study)
    elif ext == "docx":
        export_docx(path, analysis, lang, study)
    else:
        export_html(path if ext in ("html", "htm") else path + ".html", analysis, lang, study)
